#!/usr/bin/env python3
"""
智能知识库 - 文档解析器
将 PDF/DOCX/DOC/WPS/OFD/MD/TXT 解析为纯文本 Markdown
支持增量缓存（scan-cache.json），按 mtime+md5 判断是否有变更

用法:
    python3 doc-parser.py <documents_dir> <raw_output_dir>
    python3 doc-parser.py <documents_dir> <raw_output_dir> --force   # 强制重新解析
"""
import sys, os, json, hashlib, pathlib, argparse
from datetime import datetime

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

CACHE_FILE = None
SUPPORTED_EXT = {".pdf", ".docx", ".doc", ".wps", ".ofd", ".md", ".txt"}


# ==================== 文本提取 ====================

def parse_pdf(path):
    reader = PdfReader(str(path))
    texts = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            texts.append(t)
    return "\n\n".join(texts)


def parse_docx(path):
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text for c in row.cells if c.text.strip())
            if row_text.strip():
                parts.append(row_text)
    return "\n\n".join(parts)


def parse_via_libreoffice(path):
    """DOC / WPS / OFD 通过 LibreOffice 转 PDF 再解析"""
    import subprocess, tempfile
    tmpdir = tempfile.mkdtemp(prefix="kb-lo-")
    try:
        result = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", tmpdir, str(path)],
            check=False, timeout=120, capture_output=True, text=True
        )
        basename = path.stem
        pdf_path = os.path.join(tmpdir, f"{basename}.pdf")
        if os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0:
            return parse_pdf(pathlib.Path(pdf_path))

        # 备选：直接读为纯文本（某些旧格式）
        try:
            return pathlib.Path(path).read_text(encoding="utf-8", errors="replace")
        except Exception:
            pass

        raise RuntimeError(f"LibreOffice 转换失败: {result.stderr[:200] if result.stderr else '未知错误'}")
    except subprocess.TimeoutExpired:
        raise RuntimeError("LibreOffice 转换超时(120s)")
    finally:
        import shutil
        shutil.rmtree(tmpdir, ignore_errors=True)


def parse_plain(path):
    content = pathlib.Path(path).read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() == ".md" and content.startswith("---"):
        parts = content.split("---", 2)
        if len(parts) >= 3:
            content = parts[2].strip()
    return content


PARSERS = {
    ".pdf":  parse_pdf,
    ".docx": parse_docx,
    ".doc":  parse_via_libreoffice,
    ".wps":  parse_via_libreoffice,
    ".ofd":  parse_via_libreoffice,
    ".md":   parse_plain,
    ".txt":  parse_plain,
}


# ==================== 增量缓存 ====================

def load_cache():
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_cache(cache):
    os.makedirs(os.path.dirname(CACHE_FILE), exist_ok=True)
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


def file_md5(path):
    m = hashlib.md5()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(8192), b""):
            m.update(chunk)
    return m.hexdigest()


def needs_process(filepath, cache):
    key = str(filepath)
    if key not in cache:
        return True
    entry = cache[key]
    try:
        if os.path.getmtime(filepath) != entry.get("mtime", 0):
            return True
        if file_md5(filepath) != entry.get("md5", ""):
            return True
    except OSError:
        return True
    return False


# ==================== 核心处理 ====================

def extract_title(text, fallback):
    for line in text.split("\n"):
        line = line.strip()
        if line.startswith("# "):
            return line[2:].strip()
    for line in text.split("\n"):
        line = line.strip()
        if line and len(line) > 2:
            return line[:100]
    return fallback


def process_file(file_path, output_dir, documents_dir, cache, force=False):
    ext = file_path.suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        return None, f"不支持格式: {ext}"

    key = str(file_path)
    if not force and not needs_process(file_path, cache):
        return "SKIP", cache[key].get("output", "")

    try:
        text = parser(file_path)
    except Exception as e:
        return None, f"解析失败: {e}"

    if not text or not text.strip():
        return None, "文档内容为空"

    title = extract_title(text, file_path.stem)

    markdown = f"---\ntitle: {title}\nsource: {file_path.name}\n---\n\n# {title}\n\n{text}"

    try:
        rel_path = file_path.relative_to(documents_dir)
    except ValueError:
        rel_path = file_path

    out_file = pathlib.Path(output_dir) / rel_path.with_suffix(".md")
    out_file.parent.mkdir(parents=True, exist_ok=True)
    out_file.write_text(markdown, encoding="utf-8")

    cache[key] = {
        "mtime": os.path.getmtime(str(file_path)),
        "md5": file_md5(str(file_path)),
        "output": str(out_file),
        "parsed_at": datetime.now().isoformat()
    }
    return str(out_file), title


# ==================== 入口 ====================

def main():
    global CACHE_FILE

    ap = argparse.ArgumentParser(description="智能知识库 - 文档解析器")
    ap.add_argument("documents_dir")
    ap.add_argument("raw_output_dir")
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--cache-dir", default=None)
    args = ap.parse_args()

    docs = pathlib.Path(args.documents_dir).resolve()
    raw = pathlib.Path(args.raw_output_dir).resolve()
    cache_dir = pathlib.Path(args.cache_dir) if args.cache_dir else raw.parent / "db"
    CACHE_FILE = str(cache_dir / "scan-cache.json")

    if not docs.exists():
        print(f"错误: 文档目录不存在: {docs}")
        sys.exit(1)

    raw.mkdir(parents=True, exist_ok=True)
    cache_dir.mkdir(parents=True, exist_ok=True)

    cache = load_cache() if not args.force else {}
    total = success = skipped = failed = 0

    print("=" * 55)
    print("  智能知识库 - 文档解析器")
    print(f"  源目录:   {docs}")
    print(f"  输出目录: {raw}")
    print(f"  缓存文件: {CACHE_FILE}")
    print("=" * 55)

    for fpath in sorted(docs.rglob("*")):
        if not fpath.is_file():
            continue
        if fpath.suffix.lower() not in SUPPORTED_EXT:
            continue
        if fpath.name.startswith(".") or fpath.name.startswith("~"):
            continue

        total += 1
        result, info = process_file(fpath, raw, docs, cache, force=args.force)

        if result == "SKIP":
            skipped += 1
        elif result:
            success += 1
            print(f"  ✅ {fpath.relative_to(docs)}")
        else:
            failed += 1
            print(f"  ❌ {fpath.relative_to(docs)} — {info}")

    save_cache(cache)

    print("-" * 55)
    print(f"  解析完成: 共 {total} 个文件")
    print(f"  成功: {success}  跳过(未变更): {skipped}  失败: {failed}")
    print("=" * 55)


if __name__ == "__main__":
    main()
