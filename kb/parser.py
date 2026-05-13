from __future__ import annotations

import os
import subprocess
import tempfile
from pathlib import Path
from typing import List

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None

try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None

SUPPORTED_EXT = {".pdf", ".docx", ".doc", ".wps", ".ofd", ".md", ".txt"}


def parse_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext in {".doc", ".wps", ".ofd"}:
        return parse_via_libreoffice(path)
    if ext in {".md", ".txt"}:
        return parse_plain(path)
    raise ValueError(f"不支持的文档格式: {ext}")


def parse_plain(path: Path) -> str:
    for enc in ["utf-8", "utf-8-sig", "gbk", "gb18030"]:
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def parse_pdf(path: Path) -> str:
    if PdfReader is None:
        raise RuntimeError("缺少 pypdf 依赖")
    reader = PdfReader(str(path))
    texts: List[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            texts.append(text.strip())
    return "\n\n".join(texts)


def parse_docx(path: Path) -> str:
    if Document is None:
        raise RuntimeError("缺少 python-docx 依赖")
    doc = Document(str(path))
    parts: List[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def parse_via_libreoffice(path: Path) -> str:
    """DOC/WPS/OFD 通过 LibreOffice 转 PDF 再解析。"""
    with tempfile.TemporaryDirectory(prefix="kb-lo-") as tmp:
        result = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(path)],
            capture_output=True,
            text=True,
            timeout=180,
        )
        pdf = Path(tmp) / f"{path.stem}.pdf"
        if not pdf.exists() or pdf.stat().st_size == 0:
            try:
                return parse_plain(path)
            except Exception:
                raise RuntimeError(f"LibreOffice 转换失败: {result.stderr[:300]}")
        return parse_pdf(pdf)
